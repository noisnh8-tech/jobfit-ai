# -*- coding: utf-8 -*-
"""
resume_input/portfolio_pptx.py — 개인 포트폴리오(.pptx) JD 맞춤 재배치
=====================================================================
[공개 저장소 방침(2026-09-06 확정): "완성된 개인 문서 파일을 저장소에서
 직접 다운로드할 수 있게 하지는 않는다"는 방침에 따라, 원본 포트폴리오
 (assets/demo/portfolio_source.pptx)는 .gitignore 대상이라 공개
 저장소에는 없다. 이 파일(재배치 로직, _transform 이하)은 완전히
 범용이고 개인 정보를 담고 있지 않다 - 원본 .pptx가 로컬에 있을 때만
 동작하고, 없으면 generate()가 status="unavailable"을 돌려줘 호출부가
 안내 문구만 보여준다. 로컬에서 본인 포트폴리오로 이 기능을 쓰려면
 assets/demo/portfolio_source.pptx 자리에 본인 파일을 놓으면 된다
 (구조는 12장 고정 - 아래 _PROJECT_SLIDE_IDX 등 참고).]

S4(지원 준비)에서, **원본 .pptx가 로컬에 있을 때만** "맞춤 포트폴리오
다운로드" 버튼을 띄우고 이 모듈이 만든 .pptx 를 내려준다.

입력은 이미 계산돼 있는 quick_analysis 의 project_priority 하나뿐이다.
LLM 호출 없음. 문장은 생성·수정·요약하지 않는다 — 순서·번호만 바꾼다.

포폴 고정 구조(12장, 2026-09-05 "JD추가"본 - 프로젝트 3개로 확장):
  0 Cover / 1 About(3-project 카드) / 2‑4 JOBFIT AI(WHY·HOW·RESULT) /
  5‑7 Airbnb / 8‑10 Starbucks Rewards / 11 Thank you

project_priority(quick_analysis.resume_order.project_priority, 실제 이력서
프로젝트명 그대로) 순서대로 3개 프로젝트 블록을 재배치한다:
  ① About 카드: 좌·중·우 슬롯(01/02/03) 위치·뱃지는 고정, 내용(제목/설명/
     MY ROLE/팀·기간)만 새 순서의 프로젝트 것으로 교체
  ② 슬라이드 블록: 프로젝트별 3장(WHY/HOW/RESULT)을 통째로 새 순서로 재배치
  ③ 각 프로젝트 헤더의 "PROJECT 0N"을 새 위치 번호로 교체(프로젝트명은
     그대로 - 블록 전체가 같이 이동하므로 자동으로 맞음)
  ④ 우하단 페이지번호: 새 위치에 맞춰 재부여
기본 순서(JOBFIT→Airbnb→Starbucks)와 같으면 원본을 그대로 복사해 돌려준다.

반환: {"status": "ok", "pptx_path": ...} | {"status": "error", "reason": ...}
error 면 호출부는 포트폴리오 버튼을 그냥 감춘다(이력서 다운로드에는 영향 없음).
"""
from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
_CACHE_DIR = _ROOT / "data" / "portfolio_pptx_cache"
# 로컬 전용 원본 경로 - .gitignore 대상이라 공개 저장소에는 없다. 본인
# 포트폴리오로 쓰려면 이 경로에 파일을 놓거나, 이 리스트 맨 앞에 다른
# 경로를 추가하면 된다(있으면 그게 우선된다).
_SRC_CANDIDATES: list[Path] = [
    _ROOT / "assets" / "demo" / "portfolio_source.pptx",
]

# 내용/로직 버전 — 바꾸면 캐시 무효화.
_VERSION = "portfolio-pptx-v2-20260905"

# 포폴 고정 슬라이드 인덱스(원본 파일 기준, 0-base)
_COVER_IDX = 0
_ABOUT_IDX = 1
_THANKS_IDX = 11
_TOTAL_SLIDES = 12

_PROJECT_SLIDE_IDX = {
    "jobfit": (2, 3, 4),
    "airbnb": (5, 6, 7),
    "starbucks": (8, 9, 10),
}
_DEFAULT_ORDER = ["jobfit", "airbnb", "starbucks"]
_PROJECT_ORIG_NUMBER = {"jobfit": 1, "airbnb": 2, "starbucks": 3}

# project_priority(이력서 실제 프로젝트명)에서 어느 프로젝트인지 판별할 키워드.
_PROJECT_KEYS = {
    "jobfit": ("jobfit", "취업 의사결정", "ai 기반 취업", "ai 기 반 취업"),
    "airbnb": ("airbnb", "에어비앤비", "에어bnb", "숙소"),
    "starbucks": ("리워드", "starbucks", "스타벅스", "reward"),
}


def _src_pptx() -> Path | None:
    for p in _SRC_CANDIDATES:
        if p.exists():
            return p
    return None


def _in(emu: int) -> float:
    return emu / 914400


def _rank(project_priority: list[str]) -> list[str] | None:
    """project_priority → ["jobfit"|"airbnb"|"starbucks" 3개, 새 순서]
    또는 None(기본 순서와 같음 - 변환 불필요)."""
    found: list[str] = []
    for name in project_priority or []:
        low = str(name).lower()
        for key in _DEFAULT_ORDER:
            if key in found:
                continue
            if any(k in low for k in _PROJECT_KEYS[key]):
                found.append(key)
                break
    for key in _DEFAULT_ORDER:
        if key not in found:
            found.append(key)
    return None if found == _DEFAULT_ORDER else found


# ─────────────────────────────────────────────────────────────
# 변환 (python-pptx)
# ─────────────────────────────────────────────────────────────
def _set_text_keep_format(tf, new: str) -> None:
    """단일 문단 박스 전용 - 첫 run 서식을 유지한 채 텍스트만 교체(나머지
    run 제거). 페이지번호처럼 문단이 항상 1개인 박스에만 쓴다."""
    para = tf.paragraphs[0]
    runs = list(para.runs)
    if not runs:
        para.add_run().text = new
        return
    runs[0].text = new
    for extra in runs[1:]:
        extra._r.getparent().remove(extra._r)


def _para_texts(tf) -> list[str]:
    return [p.text for p in tf.paragraphs]


def _set_paragraph_texts(tf, texts: list[str]) -> None:
    """tf 의 문단 수를 texts 길이에 맞추면서(부족하면 마지막 문단을 복제해
    추가, 남으면 뒤에서 제거) 문단별 텍스트를 교체한다(문단마다 첫 run
    서식 유지, 나머지 run 제거). About 카드 설명처럼 2줄(2개 문단)인
    필드가 있어서, 단순히 .text 문자열만 통째로 옮기면 대상 박스의 옛
    두 번째 줄이 안 지워지고 남는 문제가 있었다(2026-09-05 발견)."""
    import copy

    paras = list(tf.paragraphs)
    while len(paras) < len(texts):
        new_p = copy.deepcopy(paras[-1]._p)
        paras[-1]._p.addnext(new_p)
        paras = list(tf.paragraphs)
    while len(paras) > len(texts):
        extra = paras.pop()
        extra._p.getparent().remove(extra._p)
    for para, text in zip(paras, texts):
        runs = list(para.runs)
        if not runs:
            para.add_run().text = text
            continue
        runs[0].text = text
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)


def _relabel_project_numbers(prs, order: list[str]) -> None:
    """각 프로젝트 블록(원본 인덱스 고정 상태) 안의 "PROJECT 0N" 텍스트를
    새 위치 번호로 바꾼다. 프로젝트명(JOBFIT AI 등)은 그대로 - 블록 전체가
    같이 이동하므로 손댈 필요 없다."""
    new_number = {key: i + 1 for i, key in enumerate(order)}
    for key, orig_n in _PROJECT_ORIG_NUMBER.items():
        new_n = new_number[key]
        if new_n == orig_n:
            continue
        old_txt, new_txt = f"PROJECT 0{orig_n}", f"PROJECT 0{new_n}"
        for idx in _PROJECT_SLIDE_IDX[key]:
            for sh in prs.slides[idx].shapes:
                if not sh.has_text_frame:
                    continue
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        if old_txt in r.text:
                            r.text = r.text.replace(old_txt, new_txt)


def _reorder_about_cards(prs, order: list[str]) -> None:
    """About 슬라이드의 좌/중/우 3칸 카드. 뱃지(01/02/03)·위치는 고정,
    내용(제목/한줄설명/MY ROLE 역할/팀·기간)만 새 순서의 프로젝트 것으로
    교체한다. 칸 구분은 슬라이드 폭 3등분, 칸 안 순서 구분은 top 좌표."""
    s = prs.slides[_ABOUT_IDX]
    boxes = [
        sh for sh in s.shapes
        if sh.has_text_frame and sh.text_frame.text.strip()
        and 3.0 < _in(sh.top) < 6.0
    ]
    third = _in(prs.slide_width) / 3
    cols: dict[int, list] = {0: [], 1: [], 2: []}
    for b in boxes:
        slot = 0 if _in(b.left) < third else (1 if _in(b.left) < 2 * third else 2)
        cols[slot].append(b)
    if any(len(cols[i]) != 6 for i in range(3)):
        raise RuntimeError(f"About 카드 구조가 예상과 다릅니다({[len(cols[i]) for i in range(3)]})")

    parsed: dict[int, dict] = {}
    for slot, col in cols.items():
        col_sorted = sorted(col, key=lambda b: _in(b.top))
        badge = next((b for b in col_sorted if b.text_frame.text.strip() in ("01", "02", "03")), None)
        role_label = next((b for b in col_sorted if b.text_frame.text.strip() == "MY ROLE"), None)
        if badge is None or role_label is None:
            raise RuntimeError("About 카드에서 뱃지/MY ROLE 을 찾지 못했습니다")
        rest = sorted(
            (b for b in col_sorted if b is not badge and b is not role_label),
            key=lambda b: _in(b.top),
        )
        if len(rest) != 4:
            raise RuntimeError(f"About 카드 필드 수가 예상과 다릅니다(slot={slot}, n={len(rest)})")
        title_b, desc_b, role_desc_b, team_b = rest
        # 문자열(.text)이 아니라 문단별 텍스트 리스트로 스냅샷 - 설명처럼
        # 2줄(2개 문단)인 필드가 있어 문단 수까지 그대로 옮겨야 한다.
        parsed[slot] = {
            "title": _para_texts(title_b.text_frame), "desc": _para_texts(desc_b.text_frame),
            "role": _para_texts(role_desc_b.text_frame), "team": _para_texts(team_b.text_frame),
            "boxes": (title_b, desc_b, role_desc_b, team_b),
        }

    content_by_key = {_DEFAULT_ORDER[slot]: parsed[slot] for slot in range(3)}
    for slot in range(3):
        c = content_by_key[order[slot]]
        title_b, desc_b, role_desc_b, team_b = parsed[slot]["boxes"]
        _set_paragraph_texts(title_b.text_frame, c["title"])
        _set_paragraph_texts(desc_b.text_frame, c["desc"])
        _set_paragraph_texts(role_desc_b.text_frame, c["role"])
        _set_paragraph_texts(team_b.text_frame, c["team"])


def _reorder_projects(prs, order: list[str]) -> None:
    lst = prs.slides._sldIdLst
    ids = list(lst)
    if len(ids) != _TOTAL_SLIDES:
        raise RuntimeError(f"슬라이드 수가 {_TOTAL_SLIDES}가 아닙니다({len(ids)})")
    new = [ids[_COVER_IDX], ids[_ABOUT_IDX]]
    for key in order:
        new.extend(ids[i] for i in _PROJECT_SLIDE_IDX[key])
    new.append(ids[_THANKS_IDX])
    for e in ids:
        lst.remove(e)
    for e in new:
        lst.append(e)


def _renumber_footers(prs) -> None:
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if (t.isdigit() and len(t) <= 2
                    and _in(sh.left) > 11.5 and _in(sh.top) > 6.8):
                _set_text_keep_format(sh.text_frame, str(i + 1))


def _transform(src: Path, out: Path, order: list[str]) -> None:
    from pptx import Presentation
    prs = Presentation(str(src))
    _relabel_project_numbers(prs, order)  # 슬라이드 인덱스 고정 상태에서 텍스트 먼저
    _reorder_about_cards(prs, order)
    _reorder_projects(prs, order)
    _renumber_footers(prs)                # 재배치 후 페이지번호
    prs.save(str(out))


def _export_pdf(pptx_path: Path, pdf_path: Path) -> bool:
    """PowerPoint COM 으로 pptx → pdf. 하이퍼링크는 그대로 유지된다.
    PowerPoint 미설치/COM 실패면 False(호출부는 PDF 버튼만 생략)."""
    try:
        import pythoncom
        import win32com.client
    except Exception:
        return False
    pythoncom.CoInitialize()
    pp = pres = None
    try:
        pp = win32com.client.DispatchEx("PowerPoint.Application")
        pres = pp.Presentations.Open(
            str(pptx_path), WithWindow=False, ReadOnly=True, Untitled=False)
        pres.SaveAs(str(pdf_path), 32)  # 32 = ppSaveAsPDF
        return pdf_path.exists() and pdf_path.stat().st_size > 0
    except Exception:
        return False
    finally:
        try:
            if pres is not None:
                pres.Close()
        except Exception:
            pass
        try:
            if pp is not None:
                pp.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()


# ─────────────────────────────────────────────────────────────
# 공개 엔트리
# ─────────────────────────────────────────────────────────────
def generate(project_priority: list[str] | None) -> dict:
    src = _src_pptx()
    if src is None:
        return {"status": "unavailable", "reason": "개인 문서 템플릿은 공개 저장소에서 제외되었습니다."}

    order = _rank(project_priority or [])
    order_key = "-".join(order) if order else "default"
    key = hashlib.sha1(f"{_VERSION}|{order_key}".encode("utf-8")).hexdigest()[:16]
    _CACHE_DIR.mkdir(parents=True, exist_ok=True)
    out = _CACHE_DIR / f"맞춤_포트폴리오_{order_key}_{key}.pptx"
    pdf = out.with_suffix(".pdf")

    if not (out.exists() and out.stat().st_size > 0):
        try:
            if order is None:
                # 기본 순서(JOBFIT→Airbnb→Starbucks)와 같음 → 원본 그대로.
                shutil.copy(src, out)
            else:
                _transform(src, out, order)
        except Exception as e:  # noqa: BLE001
            return {"status": "error", "reason": f"{type(e).__name__}: {e}"}

    if not (pdf.exists() and pdf.stat().st_size > 0):
        _export_pdf(out, pdf)  # 실패해도 무시(pptx 는 항상 제공)

    return {
        "status": "ok",
        "pptx_path": str(out),
        "pdf_path": str(pdf) if pdf.exists() and pdf.stat().st_size > 0 else None,
        "order": order_key,
    }
